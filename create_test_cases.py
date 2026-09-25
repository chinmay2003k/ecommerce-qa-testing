
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

file_path = "test-cases/Test_Cases.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Test Cases"

headers = [
    "TC ID",
    "Module",
    "Testing Type",
    "Test Scenario",
    "Precondition",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Priority",
    "Status"
]

ws.append(headers)

test_cases = [
    [
        "TC001", "Login", "Smoke",
        "Verify login with valid credentials",
        "Application is running",
        "1. Open login page\n2. Enter username\n3. Enter password\n4. Click Login",
        "testuser / Test@123",
        "User should login successfully and products page should open",
        "High", "Not Executed"
    ],
    [
        "TC002", "Login", "Negative",
        "Verify login with invalid password",
        "Login page is open",
        "1. Enter valid username\n2. Enter invalid password\n3. Click Login",
        "testuser / Wrong@123",
        "Error message should be displayed",
        "High", "Not Executed"
    ],
    [
        "TC003", "Login", "Negative",
        "Verify login with empty username",
        "Login page is open",
        "1. Leave username empty\n2. Enter password\n3. Click Login",
        "Password: Test@123",
        "Login should not be successful",
        "Medium", "Not Executed"
    ],
    [
        "TC004", "Login", "Negative",
        "Verify login with empty password",
        "Login page is open",
        "1. Enter username\n2. Leave password empty\n3. Click Login",
        "Username: testuser",
        "Login should not be successful",
        "Medium", "Not Executed"
    ],
    [
        "TC005", "Login", "Functional",
        "Verify logout functionality",
        "User is logged in",
        "1. Click Logout",
        "Logged-in user",
        "User should be logged out and redirected to login page",
        "Medium", "Not Executed"
    ],
    [
        "TC006", "Products", "Smoke",
        "Verify products are displayed",
        "User is logged in",
        "1. Open products page",
        "Valid user",
        "All available products should be displayed",
        "High", "Not Executed"
    ],
    [
        "TC007", "Products", "Functional",
        "Verify product name and price",
        "Products page is open",
        "1. Check product details",
        "Laptop",
        "Product name and price should be displayed correctly",
        "Medium", "Not Executed"
    ],
    [
        "TC008", "Products", "Functional",
        "Add a product to cart",
        "User is logged in",
        "1. Select a product\n2. Click Add to Cart",
        "Laptop",
        "Selected product should be added to cart",
        "High", "Not Executed"
    ],
    [
        "TC009", "Products", "Functional",
        "Add multiple products to cart",
        "User is logged in",
        "1. Add Laptop\n2. Add Mouse\n3. Add Keyboard",
        "Multiple products",
        "All selected products should appear in cart",
        "High", "Not Executed"
    ],
    [
        "TC010", "Products", "Regression",
        "Verify products page after adding product",
        "Product has been added",
        "1. Add product to cart\n2. Return to products page",
        "Laptop",
        "Products page should continue to work correctly",
        "Medium", "Not Executed"
    ],
    [
        "TC011", "Cart", "Smoke",
        "Verify cart page opens",
        "User is logged in",
        "1. Add a product\n2. Open View Cart",
        "Laptop",
        "Cart page should open successfully",
        "High", "Not Executed"
    ],
    [
        "TC012", "Cart", "Functional",
        "Verify added product appears in cart",
        "Product is added to cart",
        "1. Open cart",
        "Laptop",
        "Added product should be displayed",
        "High", "Not Executed"
    ],
    [
        "TC013", "Cart", "Functional",
        "Verify total price calculation",
        "Product is added to cart",
        "1. Add product\n2. Open cart\n3. Check total",
        "Laptop ₹55000",
        "Cart total should match product price",
        "High", "Not Executed"
    ],
    [
        "TC014", "Cart", "Functional",
        "Verify total price for multiple products",
        "Multiple products are added",
        "1. Add Laptop\n2. Add Mouse\n3. Open cart",
        "Laptop + Mouse",
        "Total should equal sum of all product prices",
        "High", "Not Executed"
    ],
    [
        "TC015", "Cart", "Regression",
        "Verify cart data before checkout",
        "Products are added to cart",
        "1. Open cart\n2. Verify products and total",
        "Multiple products",
        "Cart data should be displayed correctly",
        "Medium", "Not Executed"
    ],
    [
        "TC016", "Checkout", "Smoke",
        "Verify checkout page opens",
        "Cart contains product",
        "1. Open cart\n2. Click Proceed to Checkout",
        "Laptop",
        "Checkout page should open",
        "High", "Not Executed"
    ],
    [
        "TC017", "Checkout", "Functional",
        "Verify checkout with valid customer details",
        "Checkout page is open",
        "1. Enter name\n2. Enter address\n3. Select payment method\n4. Place order",
        "Chinmay / Mumbai / UPI",
        "Order should be placed successfully",
        "High", "Not Executed"
    ],
    [
        "TC018", "Checkout", "Negative",
        "Verify checkout with empty customer name",
        "Checkout page is open",
        "1. Leave name empty\n2. Enter address\n3. Select payment\n4. Place order",
        "Name: Empty",
        "Order should not be placed",
        "High", "Not Executed"
    ],
    [
        "TC019", "Checkout", "Negative",
        "Verify checkout with empty address",
        "Checkout page is open",
        "1. Enter name\n2. Leave address empty\n3. Select payment\n4. Place order",
        "Address: Empty",
        "Order should not be placed",
        "High", "Not Executed"
    ],
    [
        "TC020", "Checkout", "Negative",
        "Verify checkout without payment method",
        "Checkout page is open",
        "1. Enter name\n2. Enter address\n3. Do not select payment\n4. Place order",
        "Payment: Empty",
        "Order should not be placed",
        "High", "Not Executed"
    ],
    [
        "TC021", "Checkout", "Functional",
        "Verify order confirmation message",
        "Valid checkout details are entered",
        "1. Complete checkout\n2. Place order",
        "Valid customer details",
        "Order placed successfully message should be displayed",
        "High", "Not Executed"
    ],
    [
        "TC022", "Checkout", "Regression",
        "Verify cart is cleared after successful order",
        "Order is successfully placed",
        "1. Place an order\n2. Return to cart",
        "Valid order",
        "Cart should be empty after successful order",
        "Medium", "Not Executed"
    ],
    [
        "TC023", "Application", "Smoke",
        "Verify application opens successfully",
        "Application is running",
        "1. Open application URL",
        "http://127.0.0.1:5000",
        "Login page should be displayed",
        "High", "Not Executed"
    ],
    [
        "TC024", "Application", "Retesting",
        "Retest login after login defect fix",
        "Login defect has been fixed",
        "1. Enter valid credentials\n2. Click Login",
        "testuser / Test@123",
        "User should login successfully",
        "High", "Not Executed"
    ],
    [
        "TC025", "Application", "Regression",
        "Verify complete purchase flow",
        "Application is available",
        "1. Login\n2. Select product\n3. Add to cart\n4. Checkout\n5. Place order",
        "Valid user and product",
        "Complete purchase flow should work successfully",
        "High", "Not Executed"
    ]
]

for row in test_cases:
    ws.append(row)

# Header formatting
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Wrap text
for row in ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

# Column widths
widths = {
    "A": 10,
    "B": 15,
    "C": 15,
    "D": 35,
    "E": 25,
    "F": 45,
    "G": 25,
    "H": 50,
    "I": 12,
    "J": 15
}

for column, width in widths.items():
    ws.column_dimensions[column].width = width

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

wb.save(file_path)

print(f"Test case file created successfully: {file_path}")
print(f"Total test cases: {len(test_cases)}")
