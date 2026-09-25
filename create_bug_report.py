
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Bug Reports"

# Headers
headers = [
    "Bug ID",
    "Module",
    "Bug Title",
    "Severity",
    "Priority",
    "Precondition",
    "Steps to Reproduce",
    "Test Data",
    "Expected Result",
    "Actual Result",
    "Status"
]

ws.append(headers)

# Bug 1
ws.append([
    "BUG-001",
    "Cart",
    "Same product is added as duplicate cart entries",
    "Medium",
    "Medium",
    "User is logged in and is on Products page",
    "1. Login with valid credentials\n"
    "2. Add the same product twice\n"
    "3. Open Cart",
    "Product: Laptop",
    "Same product should be handled using quantity or should not create duplicate entries.",
    "The same product appears as two separate cart entries and the total becomes ₹110000.",
    "Open",
])

# Bug 2
ws.append([
    "BUG-002",
    "Checkout",
    "Order can be placed without selecting payment method",
    "High",
    "High",
    "User is logged in and has at least one product in cart",
    "1. Add a product to cart\n"
    "2. Proceed to Checkout\n"
    "3. Enter Full Name and Address\n"
    "4. Leave Payment Method as 'Select Payment Method'\n"
    "5. Click Place Order",
    "Payment Method: Not selected",
    "Payment method should be mandatory before placing an order.",
    "Order is successfully placed even though no payment method was selected.",
    "Open",
])

# Formatting
header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )
    cell.border = thin_border

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )
        cell.border = thin_border

# Column widths
widths = {
    "A": 12,
    "B": 15,
    "C": 40,
    "D": 12,
    "E": 12,
    "F": 35,
    "G": 55,
    "H": 25,
    "I": 45,
    "J": 45,
    "K": 12
}

for column, width in widths.items():
    ws.column_dimensions[column].width = width

ws.freeze_panes = "A2"

# Save file
file_path = "bug-reports/Bug_Report.xlsx"
wb.save(file_path)

print("Bug report created successfully!")
print("File:", file_path)
print("Total bugs:", ws.max_row - 1)
